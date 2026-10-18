import os
import shutil
import sqlite3
from pathlib import Path
from datetime import datetime, date
from tkinter import filedialog, messagebox

try:
    import openpyxl
except ImportError:
    openpyxl = None

from db.config import DB_NAME
from db.database import create_tables

PROJECT_ROOT = Path(__file__).resolve().parents[2]
DB_PATH = PROJECT_ROOT / DB_NAME

HEADER_MAP = {
    "title": ["title", "book_title", "book title", "book_name", "book name", "book", "name"],
    "first_name": ["first_name", "first name", "firstname", "first", "author_first", "author first", "author_first_name"],
    "last_name": ["last_name", "last name", "lastname", "last", "author_last", "author last", "author_last_name", "surname"],
    "author": ["author", "author_name", "author name", "author(s)", "authors", "writer", "written by", "creator"],
    "series": ["series", "series_title", "series title", "series_name", "series name"],
    "num_in_series": [
        "num_in_series", "num in series", "number_in_series", "number in series",
        "book_#", "book #", "book_num", "book num", "book_number", "book number",
        "series_#", "series #", "series_order", "series order", "volume", "vol", "vol.",
        "book no", "book_no", "book no.", "no_in_series", "no in series", "book#"
    ],
    "rating": ["rating", "score", "stars", "rate", "grade"],
    "genre": ["genre", "genre_title", "genre title", "genre_name", "genre name", "category"],
    "color_scheme": ["color_scheme", "color scheme", "color", "colour", "hex", "genre_color", "genre color", "color_code"],
    "topics": ["topics", "topic", "topic_name", "topic name", "tags", "tag", "categories", "category_list", "subjects", "subject"],
    "isbn": ["isbn", "isbn10", "isbn13", "isbn_10", "isbn_13", "isbn #", "isbn number"],
    "nls_order": ["nls_order", "nls order", "nls", "nls_#", "nls #", "nls_num", "nls number", "nls order #"],
    "description": ["description", "desc", "summary", "synopsis", "notes", "about", "overview", "details"],
    "book_read": ["book_read", "book read", "read", "read?", "is_read", "is read", "status", "completed", "finished", "have_read"],
    "published_date": ["published_date", "published date", "pub_date", "publication_date", "publication date", "published", "date", "year", "pub_year", "release_date"],
    "author_id": ["author_id", "author id"],
    "genre_id": ["genre_id", "genre id"],
    "series_id": ["series_id", "series id"],
    "topic_id": ["topic_id", "topic id"],
    "book_id": ["book_id", "book id", "book_#_id"],
    "id": ["id", "pk"]
}


def normalize_header(raw_header):
    if raw_header is None:
        return ""
    s = str(raw_header).strip().lower()
    for field, aliases in HEADER_MAP.items():
        if s in aliases or s.replace("_", " ") in aliases or s.replace(" ", "_") in aliases:
            return field
    return s


def parse_author_name(name_str):
    if not name_str:
        return "Unknown", "Author"
    name_str = str(name_str).strip()
    if not name_str:
        return "Unknown", "Author"

    if "," in name_str:
        parts = [p.strip() for p in name_str.split(",", 1)]
        last_name = parts[0]
        first_name = parts[1] if len(parts) > 1 and parts[1] else ""
        return first_name, last_name
    else:
        parts = name_str.split()
        if len(parts) == 1:
            return parts[0], ""
        else:
            first_name = " ".join(parts[:-1])
            last_name = parts[-1]
            return first_name, last_name


def parse_int_safe(val):
    if val is None or val == "":
        return None
    try:
        if isinstance(val, bool):
            return None
        if isinstance(val, (int, float)):
            return int(round(val))
        s = str(val).strip()
        if not s:
            return None
        try:
            return int(round(float(s)))
        except ValueError:
            import re
            m = re.search(r'\d+', s)
            if m:
                return int(m.group(0))
            return None
    except (ValueError, TypeError):
        return None


def parse_str_safe(val):
    if val is None:
        return None
    s = str(val).strip()
    return s if s else None


def parse_bool_safe(val):
    if val is None or val == "":
        return False
    if isinstance(val, bool):
        return val
    if isinstance(val, (int, float)):
        return bool(val)
    s = str(val).strip().lower()
    return s in ("1", "true", "t", "yes", "y", "read", "done", "checked")


def parse_date_safe(val):
    if val is None or val == "":
        return None
    if isinstance(val, (datetime, date)):
        return val.strftime("%Y-%m-%d")
    s = str(val).strip()
    return s if s else None


def parse_isbn_safe(val):
    if val is None or val == "":
        return None
    if isinstance(val, float):
        if val.is_integer():
            return str(int(val))
    return str(val).strip()


def find_or_create_author(cursor, first_name, last_name, cache=None):
    first_name = str(first_name).strip() if first_name else ""
    last_name = str(last_name).strip() if last_name else ""
    if not first_name and not last_name:
        first_name = "Unknown"
        last_name = "Author"

    key = (first_name.lower(), last_name.lower())
    if cache is not None and key in cache:
        return cache[key], False

    cursor.execute(
        "SELECT id FROM authors WHERE LOWER(TRIM(first_name)) = ? AND LOWER(TRIM(last_name)) = ?",
        (first_name.lower(), last_name.lower())
    )
    row = cursor.fetchone()
    if row:
        author_id = row[0]
        created = False
    else:
        cursor.execute(
            "INSERT INTO authors (first_name, last_name) VALUES (?, ?)",
            (first_name, last_name)
        )
        author_id = cursor.lastrowid
        created = True

    if cache is not None:
        cache[key] = author_id
    return author_id, created


def find_or_create_genre(cursor, title="None", color_scheme=None, cache=None):
    title = str(title).strip() if title else "None"
    if not title:
        title = "None"

    key = title.lower()
    if cache is not None and key in cache:
        return cache[key], False

    cursor.execute(
        "SELECT id, color_scheme FROM genre WHERE LOWER(TRIM(title)) = ?",
        (key,)
    )
    row = cursor.fetchone()
    if row:
        genre_id = row[0]
        created = False
        if color_scheme and not row[1]:
            cursor.execute("UPDATE genre SET color_scheme = ? WHERE id = ?", (color_scheme, genre_id))
    else:
        cursor.execute(
            "INSERT INTO genre (title, color_scheme) VALUES (?, ?)",
            (title, color_scheme)
        )
        genre_id = cursor.lastrowid
        created = True

    if cache is not None:
        cache[key] = genre_id
    return genre_id, created


def find_or_create_series(cursor, title, author_id, cache=None):
    title = str(title).strip() if title else ""
    if not title:
        return None, False

    key = (title.lower(), author_id)
    if cache is not None and key in cache:
        return cache[key], False

    cursor.execute(
        "SELECT id FROM series WHERE LOWER(TRIM(title)) = ? AND author_id = ?",
        (title.lower(), author_id)
    )
    row = cursor.fetchone()
    if row:
        series_id = row[0]
        created = False
    else:
        cursor.execute(
            "SELECT id, author_id FROM series WHERE LOWER(TRIM(title)) = ?",
            (title.lower(),)
        )
        row = cursor.fetchone()
        if row:
            series_id = row[0]
            if row[1] is None and author_id is not None:
                cursor.execute("UPDATE series SET author_id = ? WHERE id = ?", (author_id, series_id))
            created = False
        else:
            cursor.execute(
                "INSERT INTO series (title, author_id) VALUES (?, ?)",
                (title, author_id)
            )
            series_id = cursor.lastrowid
            created = True

    if cache is not None:
        cache[key] = series_id
    return series_id, created


def find_or_create_topic(cursor, name, cache=None):
    name = str(name).strip() if name else ""
    if not name:
        return None, False

    key = name.lower()
    if cache is not None and key in cache:
        return cache[key], False

    cursor.execute(
        "SELECT id FROM topics WHERE LOWER(TRIM(name)) = ?",
        (name.lower(),)
    )
    row = cursor.fetchone()
    if row:
        topic_id = row[0]
        created = False
    else:
        cursor.execute(
            "INSERT INTO topics (name) VALUES (?)",
            (name,)
        )
        topic_id = cursor.lastrowid
        created = True

    if cache is not None:
        cache[key] = topic_id
    return topic_id, created


def link_book_topic(cursor, book_id, topic_id):
    if not book_id or not topic_id:
        return False
    cursor.execute(
        "SELECT id FROM book_topics WHERE book_id = ? AND topic_id = ?",
        (book_id, topic_id)
    )
    if not cursor.fetchone():
        cursor.execute(
            "INSERT INTO book_topics (book_id, topic_id) VALUES (?, ?)",
            (book_id, topic_id)
        )
        return True
    return False


def extract_sheet_rows(sheet):
    """
    Extracts headers and list of row dicts from an openpyxl sheet.
    """
    rows_iter = sheet.iter_rows(values_only=True)
    header_row = None
    for row in rows_iter:
        if row and any(cell is not None and str(cell).strip() != "" for cell in row):
            header_row = row
            break

    if header_row is None:
        return []

    col_map = {}
    for idx, cell in enumerate(header_row):
        if cell is not None and str(cell).strip() != "":
            norm = normalize_header(cell)
            if norm:
                col_map[idx] = norm

    data_rows = []
    for row in rows_iter:
        if not row or not any(cell is not None and str(cell).strip() != "" for cell in row):
            continue
        row_dict = {}
        for idx, norm_name in col_map.items():
            if idx < len(row):
                row_dict[norm_name] = row[idx]
        if any(v is not None and str(v).strip() != "" for v in row_dict.values()):
            data_rows.append(row_dict)

    return data_rows


def import_sheet_positional(sheet, cursor, stats, author_cache=None, series_cache=None, genre_cache=None):
    """
    Imports a sheet following the positional column rules:
    - First row is ignored.
    - Column A: Author -> add/find in DB, link author_id.
    - Column B: Series -> add/find in DB linking author to series. If blank, skipped.
    - Column C: Number in series -> used if Column B is not empty.
    - Column D: Book title -> add book to DB linking author, series, num_in_series.
    - Genre: "None" (if "none" is not in DB, create it).
    """
    rows = list(sheet.iter_rows(values_only=True))
    if len(rows) <= 1:
        return

    # First row will be ignored
    data_rows = rows[1:]

    for row in data_rows:
        if not row or not any(cell is not None and str(cell).strip() != "" for cell in row):
            continue

        raw_author = row[0] if len(row) > 0 else None
        raw_series = row[1] if len(row) > 1 else None
        raw_num = row[2] if len(row) > 2 else None
        raw_title = row[3] if len(row) > 3 else None

        title = parse_str_safe(raw_title)
        if not title:
            # If no book title in column D, skip row
            continue

        # Column A: Author
        first_name, last_name = parse_author_name(raw_author)
        author_id, author_created = find_or_create_author(cursor, first_name, last_name, cache=author_cache)
        if author_created:
            stats["authors"] += 1

        # Column B: Series (linking author to series; if cell value is blank, skipped)
        series_title = parse_str_safe(raw_series)
        if series_title:
            series_id, series_created = find_or_create_series(cursor, series_title, author_id, cache=series_cache)
            if series_created:
                stats["series"] += 1
            # Column C: Number in series value (if Column B is not empty)
            num_in_series = parse_int_safe(raw_num)
        else:
            series_id = None
            num_in_series = None

        # Genre: "None" (if "none" is not an option for the database then create it)
        genre_id, genre_created = find_or_create_genre(cursor, "None", cache=genre_cache)
        if genre_created:
            stats["genres"] += 1

        # Add book to database
        cursor.execute(
            """INSERT INTO books (
                title, series_id, num_in_series, book_read, nls_order,
                published_date, rating, isbn, genre_id, author_id, description
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            (
                title, series_id, num_in_series, 0, None,
                None, None, None, genre_id, author_id, None
            )
        )
        stats["books"] += 1


def import_excel_data(file_path, db_path=None):
    """
    Imports an Excel workbook file into the SQLite database.
    Returns a dict with statistics on imported entities.
    """
    if openpyxl is None:
        raise ImportError("openpyxl is not installed. Please install openpyxl to import Excel files.")

    target_db = Path(db_path) if db_path else DB_PATH
    create_tables(target_db)

    wb = openpyxl.load_workbook(file_path, data_only=True)

    stats = {
        "books": 0,
        "authors": 0,
        "series": 0,
        "genres": 0,
        "topics": 0
    }

    conn = sqlite3.connect(target_db)
    cursor = conn.cursor()

    try:
        author_cache = {}
        genre_cache = {}
        series_cache = {}
        topic_cache = {}

        sheet_lookup = {sheet_name.strip().lower(): sheet_name for sheet_name in wb.sheetnames}

        # Check if this is a relational multi-sheet workbook export
        is_relational = False
        if "authors" in sheet_lookup and "books" in sheet_lookup and (
            "genres" in sheet_lookup or "series" in sheet_lookup or "topics" in sheet_lookup
        ):
            books_sheet = wb[sheet_lookup.get("books") or sheet_lookup.get("book")]
            rows_preview = extract_sheet_rows(books_sheet)
            if rows_preview and ("author_id" in rows_preview[0] or "genre_id" in rows_preview[0]):
                is_relational = True

        if is_relational:
            author_id_map = {}
            genre_id_map = {}
            series_id_map = {}
            topic_id_map = {}
            book_id_map = {}

            # 1. Process Authors sheet
            authors_sheet_name = sheet_lookup.get("authors") or sheet_lookup.get("author")
            if authors_sheet_name:
                sheet = wb[authors_sheet_name]
                for row in extract_sheet_rows(sheet):
                    excel_id = row.get("id") or row.get("author_id")
                    first_name = row.get("first_name")
                    last_name = row.get("last_name")
                    if not first_name and not last_name and row.get("author"):
                        first_name, last_name = parse_author_name(row.get("author"))

                    if first_name or last_name:
                        author_id, created = find_or_create_author(cursor, first_name, last_name, cache=author_cache)
                        if created:
                            stats["authors"] += 1
                        if excel_id is not None:
                            author_id_map[excel_id] = author_id

            # 2. Process Genres sheet
            genres_sheet_name = sheet_lookup.get("genres") or sheet_lookup.get("genre")
            if genres_sheet_name:
                sheet = wb[genres_sheet_name]
                for row in extract_sheet_rows(sheet):
                    excel_id = row.get("id") or row.get("genre_id")
                    title = row.get("genre") or row.get("title") or row.get("name")
                    color = row.get("color_scheme") or row.get("color")
                    if title:
                        genre_id, created = find_or_create_genre(cursor, title, color_scheme=color, cache=genre_cache)
                        if created:
                            stats["genres"] += 1
                        if excel_id is not None:
                            genre_id_map[excel_id] = genre_id

            # 3. Process Series sheet
            series_sheet_name = sheet_lookup.get("series")
            if series_sheet_name:
                sheet = wb[series_sheet_name]
                for row in extract_sheet_rows(sheet):
                    excel_id = row.get("id") or row.get("series_id")
                    title = row.get("series") or row.get("title") or row.get("name")
                    author_ref = row.get("author_id")
                    if author_ref in author_id_map:
                        author_id = author_id_map[author_ref]
                    elif row.get("author") or row.get("first_name") or row.get("last_name"):
                        first_name = row.get("first_name")
                        last_name = row.get("last_name")
                        if not first_name and not last_name:
                            first_name, last_name = parse_author_name(row.get("author"))
                        author_id, created = find_or_create_author(cursor, first_name, last_name, cache=author_cache)
                        if created:
                            stats["authors"] += 1
                    else:
                        author_id, created = find_or_create_author(cursor, "Unknown", "Author", cache=author_cache)
                        if created:
                            stats["authors"] += 1

                    if title:
                        series_id, created = find_or_create_series(cursor, title, author_id, cache=series_cache)
                        if created:
                            stats["series"] += 1
                        if excel_id is not None:
                            series_id_map[excel_id] = series_id

            # 4. Process Topics sheet
            topics_sheet_name = sheet_lookup.get("topics") or sheet_lookup.get("topic")
            if topics_sheet_name:
                sheet = wb[topics_sheet_name]
                for row in extract_sheet_rows(sheet):
                    excel_id = row.get("id") or row.get("topic_id")
                    name = row.get("topics") or row.get("topic") or row.get("name") or row.get("title")
                    if name:
                        topic_id, created = find_or_create_topic(cursor, name, cache=topic_cache)
                        if created:
                            stats["topics"] += 1
                        if excel_id is not None:
                            topic_id_map[excel_id] = topic_id

            # 5. Process Books sheet
            books_sheet_name = sheet_lookup.get("books") or sheet_lookup.get("book")
            if books_sheet_name:
                sheet = wb[books_sheet_name]
                for row in extract_sheet_rows(sheet):
                    title = parse_str_safe(row.get("title"))
                    if not title:
                        continue

                    author_ref = row.get("author_id")
                    if author_ref in author_id_map:
                        author_id = author_id_map[author_ref]
                    elif row.get("first_name") or row.get("last_name"):
                        author_id, created = find_or_create_author(
                            cursor,
                            row.get("first_name", ""),
                            row.get("last_name", ""),
                            cache=author_cache
                        )
                        if created:
                            stats["authors"] += 1
                    elif row.get("author"):
                        first_name, last_name = parse_author_name(row.get("author"))
                        author_id, created = find_or_create_author(cursor, first_name, last_name, cache=author_cache)
                        if created:
                            stats["authors"] += 1
                    else:
                        author_id, created = find_or_create_author(cursor, "Unknown", "Author", cache=author_cache)
                        if created:
                            stats["authors"] += 1

                    genre_ref = row.get("genre_id")
                    if genre_ref in genre_id_map:
                        genre_id = genre_id_map[genre_ref]
                    elif row.get("genre"):
                        genre_id, created = find_or_create_genre(cursor, row.get("genre"), cache=genre_cache)
                        if created:
                            stats["genres"] += 1
                    else:
                        genre_id, created = find_or_create_genre(cursor, "None", cache=genre_cache)
                        if created:
                            stats["genres"] += 1

                    series_ref = row.get("series_id")
                    if series_ref in series_id_map:
                        series_id = series_id_map[series_ref]
                    elif row.get("series"):
                        series_id, created = find_or_create_series(cursor, row.get("series"), author_id, cache=series_cache)
                        if created:
                            stats["series"] += 1
                    else:
                        series_id = None

                    num_in_series = parse_int_safe(row.get("num_in_series"))
                    rating = parse_int_safe(row.get("rating"))
                    isbn = parse_isbn_safe(row.get("isbn"))
                    nls_order = parse_str_safe(row.get("nls_order"))
                    description = parse_str_safe(row.get("description"))
                    book_read = parse_bool_safe(row.get("book_read"))
                    published_date = parse_date_safe(row.get("published_date"))

                    cursor.execute(
                        """INSERT INTO books (
                            title, series_id, num_in_series, book_read, nls_order,
                            published_date, rating, isbn, genre_id, author_id, description
                        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                        (
                            title, series_id, num_in_series, book_read, nls_order,
                            published_date, rating, isbn, genre_id, author_id, description
                        )
                    )
                    db_book_id = cursor.lastrowid
                    stats["books"] += 1

                    excel_book_id = row.get("id") or row.get("book_id")
                    if excel_book_id is not None:
                        book_id_map[excel_book_id] = db_book_id

                    if row.get("topics"):
                        topics_raw = str(row.get("topics"))
                        for sep in [";", "|", "\n"]:
                            topics_raw = topics_raw.replace(sep, ",")
                        topic_list = [t.strip() for t in topics_raw.split(",") if t.strip()]
                        for tname in topic_list:
                            topic_id, created = find_or_create_topic(cursor, tname, cache=topic_cache)
                            if created:
                                stats["topics"] += 1
                            link_book_topic(cursor, db_book_id, topic_id)

            # 6. Process Book_Topics sheet if present
            bt_sheet_name = (
                sheet_lookup.get("book_topics")
                or sheet_lookup.get("book topics")
                or sheet_lookup.get("booktopics")
            )
            if bt_sheet_name:
                sheet = wb[bt_sheet_name]
                for row in extract_sheet_rows(sheet):
                    book_ref = row.get("book_id") or row.get("book") or row.get("title")
                    topic_ref = row.get("topic_id") or row.get("topic") or row.get("name")

                    db_book_id = book_id_map.get(book_ref)
                    if db_book_id is None and parse_int_safe(book_ref) is not None:
                        db_book_id = parse_int_safe(book_ref)
                    if db_book_id is None and book_ref:
                        cursor.execute("SELECT id FROM books WHERE LOWER(title) = ? ORDER BY id DESC LIMIT 1", (str(book_ref).strip().lower(),))
                        brow = cursor.fetchone()
                        if brow:
                            db_book_id = brow[0]

                    db_topic_id = topic_id_map.get(topic_ref)
                    if db_topic_id is None and parse_int_safe(topic_ref) is not None:
                        db_topic_id = parse_int_safe(topic_ref)
                    if db_topic_id is None and topic_ref:
                        db_topic_id, created = find_or_create_topic(cursor, topic_ref, cache=topic_cache)
                        if created:
                            stats["topics"] += 1

                    if db_book_id and db_topic_id:
                        link_book_topic(cursor, db_book_id, db_topic_id)

        else:
            # Positional column import across all sheets in the workbook:
            # Column A: Author
            # Column B: Series (skipped if blank)
            # Column C: Number in series (if Column B is not empty)
            # Column D: Book Title
            # Genre: "None" (created in DB if not existing)
            for sheet in wb.worksheets:
                import_sheet_positional(
                    sheet,
                    cursor,
                    stats,
                    author_cache=author_cache,
                    series_cache=series_cache,
                    genre_cache=genre_cache
                )

        conn.commit()
        return stats

    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def import_excel_workbook(app=None, file_path=None):
    """
    GUI / interactive handler for importing an Excel workbook.
    Prompts user to select a file (if file_path is None), backs up the database,
    imports data, shows a completion message, and refreshes the UI.
    """
    if openpyxl is None:
        messagebox.showerror(
            "Missing Dependency",
            "openpyxl is required to import Excel files.\nPlease install openpyxl (pip install openpyxl)."
        )
        return

    if not file_path:
        file_path = filedialog.askopenfilename(
            title="Choose an Excel workbook to import",
            filetypes=[
                ("Excel Workbook", "*.xlsx *.xlsm *.xltx *.xltm *.xls"),
                ("All files", "*.*")
            ]
        )

    if not file_path:
        return

    confirmed = messagebox.askyesno(
        "Import Excel Workbook?",
        f"Do you want to import data from:\n{Path(file_path).name}\n\n"
        "Your existing database will be backed up automatically first."
    )
    if not confirmed:
        return

    try:
        # Create a backup before importing
        if DB_PATH.exists():
            recovery_path = DB_PATH.with_name(f"{DB_PATH.stem}_before_import.db")
            shutil.copy2(DB_PATH, recovery_path)

        stats = import_excel_data(file_path, DB_PATH)

        summary_msg = (
            "Excel workbook imported successfully!\n\n"
            f"• Books imported: {stats['books']}\n"
            f"• New authors added: {stats['authors']}\n"
            f"• New series added: {stats['series']}\n"
            f"• New genres added: {stats['genres']}\n"
            f"• New topics added: {stats['topics']}"
        )
        messagebox.showinfo("Import Complete", summary_msg)

        if app is not None:
            if hasattr(app, "book_table") and hasattr(app.book_table, "load_books"):
                app.book_table.load_books()
            if hasattr(app, "load_authors_into_filter"):
                app.load_authors_into_filter()
            if hasattr(app, "load_genres_into_filter"):
                app.load_genres_into_filter()
            if hasattr(app, "load_topics_into_filter"):
                app.load_topics_into_filter()
            if hasattr(app, "add_window") and app.add_window is not None:
                if hasattr(app.add_window, "refresh_book_tab_authors"):
                    app.add_window.refresh_book_tab_authors()
                if hasattr(app.add_window, "refresh_book_tab_genre"):
                    app.add_window.refresh_book_tab_genre()
                if hasattr(app.add_window, "refresh_book_tab_topics"):
                    app.add_window.refresh_book_tab_topics()

    except Exception as error:
        messagebox.showerror("Import Failed", f"Failed to import Excel workbook:\n{error}")
